"""
Generates formatted Excel workbooks for financial statements using openpyxl.
"""
from io import BytesIO
from datetime import date, datetime
from typing import Dict, Any, List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter


# Brand styles
HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="7C3AED")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="64748B")
TOTAL_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
TOTAL_FONT = Font(name="Calibri", bold=True, size=10)
BODY_FONT = Font(name="Calibri", size=10)
ALT_ROW_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)
GBP_FORMAT = '£#,##0.00'


class ExcelExportService:
    """Generate branded Excel workbooks for financial statements."""

    def __init__(self, company_name: str, downloaded_by: str):
        self.company_name = company_name
        self.downloaded_by = downloaded_by

    def _setup_sheet(self, ws, title: str, period_text: str, headers: List[str]):
        """Add branded header rows and column headers to a worksheet."""
        # Title row
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.cell(row=1, column=1, value=f"Accuratea — {self.company_name}").font = TITLE_FONT

        # Period row
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        ws.cell(row=2, column=1, value=f"{title} | {period_text}").font = SUBTITLE_FONT

        # Download metadata
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
        ws.cell(
            row=3, column=1,
            value=f"Downloaded by {self.downloaded_by} on {datetime.now().strftime('%d %b %Y at %H:%M')}"
        ).font = Font(name="Calibri", size=8, color="94A3B8")

        # Column headers (row 5)
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = THIN_BORDER

        return 6  # First data row

    def _add_data_row(self, ws, row_num: int, values: list, is_total: bool = False):
        """Add a data row with formatting."""
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=value)
            cell.font = TOTAL_FONT if is_total else BODY_FONT
            cell.border = THIN_BORDER

            if is_total:
                cell.fill = TOTAL_FILL

            elif row_num % 2 == 0:
                cell.fill = ALT_ROW_FILL

            # Currency formatting for numeric columns
            if isinstance(value, (int, float)) and col_idx > 1:
                cell.number_format = GBP_FORMAT
                cell.alignment = Alignment(horizontal="right")

    def _auto_width(self, ws, headers: List[str]):
        """Auto-fit column widths."""
        for col_idx, header in enumerate(headers, 1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(len(header) + 4, 15)


    # INCOME STATEMENT
    def generate_income_statement(self, data: Dict[str, Any], start_date: str, end_date: str) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Income Statement"

        headers = ["Account Code", "Account Name", "Debit (£)", "Credit (£)"]
        row = self._setup_sheet(ws, "Income Statement", f"{start_date} to {end_date}", headers)

        # Revenue
        ws.cell(row=row, column=1, value="REVENUE").font = Font(bold=True, color="10B981", size=10)
        row += 1

        for acc in data.get("revenue", {}).get("accounts", []):
            self._add_data_row(ws, row, [
                acc.get("account_code", ""),
                acc.get("account_name", ""),
                acc.get("debit", 0),
                acc.get("credit", 0),
            ])
            row += 1

        self._add_data_row(ws, row, [
            "", "Total Revenue", "", data.get("revenue", {}).get("total", 0)
        ], is_total=True)
        row += 2

        # Expenses
        ws.cell(row=row, column=1, value="EXPENSES").font = Font(bold=True, color="EF4444", size=10)
        row += 1

        for acc in data.get("expenses", {}).get("accounts", []):
            self._add_data_row(ws, row, [
                acc.get("account_code", ""),
                acc.get("account_name", ""),
                acc.get("debit", 0),
                acc.get("credit", 0),
            ])
            row += 1

        self._add_data_row(ws, row, [
            "", "Total Expenses", data.get("expenses", {}).get("total", 0), ""
        ], is_total=True)
        row += 2

        # Net Income
        net = data.get("net_income", 0)
        label = "Net Profit" if net >= 0 else "Net Loss"
        self._add_data_row(ws, row, ["", label, "", abs(net)], is_total=True)

        self._auto_width(ws, headers)
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


    # BALANCE SHEET
    def generate_balance_sheet(self, data: Dict[str, Any], as_of_date: str) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Balance Sheet"

        headers = ["Account Code", "Account Name", "Amount (£)"]
        row = self._setup_sheet(ws, "Balance Sheet", f"As at {as_of_date}", headers)

        for section_key, section_label in [("assets", "ASSETS"), ("liabilities", "LIABILITIES"), ("equity", "EQUITY")]:
            section = data.get(section_key, {})
            ws.cell(row=row, column=1, value=section_label).font = Font(bold=True, size=10)
            row += 1

            for acc in section.get("accounts", []):
                balance = acc.get("debit", 0) if section_key == "assets" else acc.get("credit", 0)
                self._add_data_row(ws, row, [
                    acc.get("account_code", ""),
                    acc.get("account_name", ""),
                    balance,
                ])
                row += 1

            self._add_data_row(ws, row, [
                "", f"Total {section_label.title()}", section.get("total", 0)
            ], is_total=True)
            row += 2

        self._auto_width(ws, headers)
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


    # TRIAL BALANCE
    def generate_trial_balance(self, data: Dict[str, Any], as_of_date: str) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Trial Balance"

        headers = ["Account Code", "Account Name", "Debit (£)", "Credit (£)"]
        row = self._setup_sheet(ws, "Trial Balance", f"As at {as_of_date}", headers)

        for acc in data.get("accounts", []):
            self._add_data_row(ws, row, [
                acc.get("account_code", ""),
                acc.get("account_name", ""),
                acc.get("debit", 0),
                acc.get("credit", 0),
            ])
            row += 1

        self._add_data_row(ws, row, [
            "", "TOTALS",
            data.get("total_debits", 0),
            data.get("total_credits", 0),
        ], is_total=True)

        self._auto_width(ws, headers)
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


    # ETB (landscape-style wide sheet)
    def generate_etb(self, data: Dict[str, Any], start_date: str, end_date: str) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Extended Trial Balance"

        headers = [
            "Code", "Account", "TB Dr", "TB Cr", "Adj Dr", "Adj Cr",
            "ATB Dr", "ATB Cr", "P&L Dr", "P&L Cr", "BS Dr", "BS Cr"
        ]
        row = self._setup_sheet(ws, "Extended Trial Balance", f"{start_date} to {end_date}", headers)

        for r in data.get("rows", []):
            self._add_data_row(ws, row, [
                r["account_code"], r["account_name"],
                r["tb_dr"], r["tb_cr"], r["adj_dr"], r["adj_cr"],
                r["atb_dr"], r["atb_cr"], r["pl_dr"], r["pl_cr"],
                r["bs_dr"], r["bs_cr"],
            ])
            row += 1

        totals = data.get("totals", {})
        self._add_data_row(ws, row, [
            "", "TOTALS",
            totals.get("tb_dr", 0), totals.get("tb_cr", 0),
            totals.get("adj_dr", 0), totals.get("adj_cr", 0),
            totals.get("atb_dr", 0), totals.get("atb_cr", 0),
            totals.get("pl_dr", 0), totals.get("pl_cr", 0),
            totals.get("bs_dr", 0), totals.get("bs_cr", 0),
        ], is_total=True)

        self._auto_width(ws, headers)
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


    # FULL WORKBOOK
    def generate_full_workbook(
        self,
        journal_entries: list,
        trial_balance: dict,
        income_statement: dict,
        balance_sheet: dict,
        cash_flow: dict,
        etb: dict,
        accounts: list,
        start_date: str,
        end_date: str,
    ) -> bytes:
        """Generate a single workbook with all financial statements as separate sheets."""
        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)
        period_text = f"{start_date} to {end_date}"
        as_of_text = f"As at {end_date}"
        self._add_journal_entries_sheet(wb, journal_entries, period_text)
        self._add_trial_balance_sheet(wb, trial_balance, as_of_text)
        self._add_income_statement_sheet(wb, income_statement, period_text)
        self._add_balance_sheet_sheet(wb, balance_sheet, as_of_text)
        self._add_cash_flow_sheet(wb, cash_flow, period_text)
        self._add_etb_sheet(wb, etb, period_text)
        self._add_chart_of_accounts_sheet(wb, accounts)
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


    # generate_full_workbook methods
    def _add_journal_entries_sheet(self, wb: "Workbook", entries: list, period_text: str):
        """Add journal entries as a sheet."""
        ws = wb.create_sheet("Journal Entries")
        headers = ["Date", "Entry #", "Description", "Account Code", "Account Name", "Debit (£)", "Credit (£)"]
        row = self._setup_sheet(ws, "Journal Entries", period_text, headers)

        for entry in entries:
            for line in entry.get("lines", []):
                self._add_data_row(ws, row, [
                    entry.get("date", ""),
                    entry.get("entry_number", ""),
                    entry.get("description", ""),
                    line.get("account_code", ""),
                    line.get("account_name", ""),
                    line.get("debit", 0),
                    line.get("credit", 0),
                ])
                row += 1
            # Blank row between entries
            row += 1

        self._auto_width(ws, headers)

    def _add_trial_balance_sheet(self, wb: "Workbook", data: dict, period_text: str):
        """Add trial balance as a sheet."""
        ws = wb.create_sheet("Trial Balance")
        headers = ["Account Code", "Account Name", "Debit (£)", "Credit (£)"]
        row = self._setup_sheet(ws, "Trial Balance", period_text, headers)

        for acc in data.get("accounts", []):
            self._add_data_row(ws, row, [
                acc.get("account_code", ""),
                acc.get("account_name", ""),
                acc.get("debit", 0),
                acc.get("credit", 0),
            ])
            row += 1

        self._add_data_row(ws, row, [
            "", "TOTALS",
            data.get("total_debits", 0),
            data.get("total_credits", 0),
        ], is_total=True)

        self._auto_width(ws, headers)

    def _add_income_statement_sheet(self, wb: "Workbook", data: dict, period_text: str):
        """Add income statement as a sheet."""
        ws = wb.create_sheet("Income Statement")
        headers = ["Account Code", "Account Name", "Debit (£)", "Credit (£)"]
        row = self._setup_sheet(ws, "Income Statement", period_text, headers)

        # Revenue
        ws.cell(row=row, column=1, value="REVENUE").font = Font(bold=True, color="10B981", size=10)
        row += 1
        for acc in data.get("revenue", {}).get("accounts", []):
            self._add_data_row(ws, row, [
                acc.get("account_code", ""), acc.get("account_name", ""),
                acc.get("debit", 0), acc.get("credit", 0),
            ])
            row += 1
        self._add_data_row(ws, row, [
            "", "Total Revenue", "", data.get("revenue", {}).get("total", 0)
        ], is_total=True)
        row += 2

        # Expenses
        ws.cell(row=row, column=1, value="EXPENSES").font = Font(bold=True, color="EF4444", size=10)
        row += 1
        for acc in data.get("expenses", {}).get("accounts", []):
            self._add_data_row(ws, row, [
                acc.get("account_code", ""), acc.get("account_name", ""),
                acc.get("debit", 0), acc.get("credit", 0),
            ])
            row += 1
        self._add_data_row(ws, row, [
            "", "Total Expenses", data.get("expenses", {}).get("total", 0), ""
        ], is_total=True)
        row += 2

        # Net Income
        net = data.get("net_income", 0)
        label = "Net Profit" if net >= 0 else "Net Loss"
        self._add_data_row(ws, row, ["", label, "", abs(net)], is_total=True)

        self._auto_width(ws, headers)

    def _add_balance_sheet_sheet(self, wb: "Workbook", data: dict, period_text: str):
        """Add balance sheet as a sheet."""
        ws = wb.create_sheet("Balance Sheet")
        headers = ["Account Code", "Account Name", "Amount (£)"]
        row = self._setup_sheet(ws, "Balance Sheet", period_text, headers)

        for section_key, section_label in [("assets", "ASSETS"), ("liabilities", "LIABILITIES"), ("equity", "EQUITY")]:
            section = data.get(section_key, {})
            ws.cell(row=row, column=1, value=section_label).font = Font(bold=True, size=10)
            row += 1
            for acc in section.get("accounts", []):
                balance = acc.get("debit", 0) if section_key == "assets" else acc.get("credit", 0)
                self._add_data_row(ws, row, [
                    acc.get("account_code", ""), acc.get("account_name", ""), balance,
                ])
                row += 1
            self._add_data_row(ws, row, [
                "", f"Total {section_label.title()}", section.get("total", 0)
            ], is_total=True)
            row += 2

        self._auto_width(ws, headers)

    def _add_cash_flow_sheet(self, wb: "Workbook", data: dict, period_text: str):
        """Add cash flow statement as a sheet."""
        ws = wb.create_sheet("Cash Flow")
        headers = ["Category", "Description", "Amount (£)"]
        row = self._setup_sheet(ws, "Cash Flow Statement", period_text, headers)

        for section_key, section_label in [
            ("operating", "OPERATING ACTIVITIES"),
            ("investing", "INVESTING ACTIVITIES"),
            ("financing", "FINANCING ACTIVITIES"),
        ]:
            section = data.get(section_key, {})
            ws.cell(row=row, column=1, value=section_label).font = Font(bold=True, size=10)
            row += 1
            for item in section.get("items", []):
                self._add_data_row(ws, row, [
                    "", item.get("description", ""), item.get("amount", 0),
                ])
                row += 1
            self._add_data_row(ws, row, [
                "", f"Net {section_label.title()}", section.get("total", 0)
            ], is_total=True)
            row += 2

        # Summary
        summary = data.get("summary", {})
        self._add_data_row(ws, row, ["", "Beginning Cash", summary.get("beginning_cash_balance", 0)], is_total=True)
        row += 1
        self._add_data_row(ws, row, ["", "Net Change in Cash", summary.get("net_change_in_cash", 0)], is_total=True)
        row += 1
        self._add_data_row(ws, row, ["", "Ending Cash", summary.get("ending_cash_balance", 0)], is_total=True)

        self._auto_width(ws, headers)

    def _add_etb_sheet(self, wb: "Workbook", data: dict, period_text: str):
        """Add ETB as a sheet."""
        ws = wb.create_sheet("ETB")
        headers = [
            "Code", "Account", "TB Dr", "TB Cr", "Adj Dr", "Adj Cr",
            "ATB Dr", "ATB Cr", "P&L Dr", "P&L Cr", "BS Dr", "BS Cr"
        ]
        row = self._setup_sheet(ws, "Extended Trial Balance", period_text, headers)

        for r in data.get("rows", []):
            self._add_data_row(ws, row, [
                r["account_code"], r["account_name"],
                r["tb_dr"], r["tb_cr"], r["adj_dr"], r["adj_cr"],
                r["atb_dr"], r["atb_cr"], r["pl_dr"], r["pl_cr"],
                r["bs_dr"], r["bs_cr"],
            ])
            row += 1

        totals = data.get("totals", {})
        self._add_data_row(ws, row, [
            "", "TOTALS",
            totals.get("tb_dr", 0), totals.get("tb_cr", 0),
            totals.get("adj_dr", 0), totals.get("adj_cr", 0),
            totals.get("atb_dr", 0), totals.get("atb_cr", 0),
            totals.get("pl_dr", 0), totals.get("pl_cr", 0),
            totals.get("bs_dr", 0), totals.get("bs_cr", 0),
        ], is_total=True)

        self._auto_width(ws, headers)

    def _add_chart_of_accounts_sheet(self, wb: "Workbook", accounts: list):
        """Add chart of accounts as a sheet."""
        ws = wb.create_sheet("Chart of Accounts")
        headers = ["Account Code", "Account Name", "Type", "Sub Type", "Normal Balance", "Current Balance (£)"]
        row = self._setup_sheet(ws, "Chart of Accounts", "Full Account Listing", headers)

        for acc in accounts:
            self._add_data_row(ws, row, [
                acc.get("account_code", ""),
                acc.get("account_name", ""),
                acc.get("account_type", ""),
                acc.get("sub_type", "") or "",
                acc.get("normal_balance", ""),
                acc.get("current_balance", 0),
            ])
            row += 1

        self._auto_width(ws, headers)